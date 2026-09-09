"""Deterministic import, prices, review state and WooCommerce export. No network."""
from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import re
import sqlite3
import unicodedata
from collections import Counter, defaultdict
from contextlib import contextmanager
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from urllib.parse import urlsplit
from uuid import uuid4


class ValidationError(ValueError):
    pass


def now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def norm(value):
    value = str(value or "").replace("đ", "d").replace("Đ", "D")
    return " ".join(re.sub(r"[^a-z0-9]+", " ", unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode().lower()).split())


FIELDS = {
    "name": "Naziv proizvoda *", "sku": "SKU / šifra artikla", "ean": "EAN / GTIN",
    "model": "Proizvođački model / MPN", "vpc": "Redovna VPC", "promo_vpc": "Akcija VPC",
    "brand": "Proizvođač / brend", "category": "Kategorija", "source_url": "Proizvođačka stranica (URL)",
    "images": "URL-ovi slika", "short": "Kratki opis", "description": "Dugi opis",
}
ALIASES = {
    "name": ["naziv", "naziv proizvoda", "naziv artikla", "proizvod", "artikl", "name", "product name", "opis artikla"],
    "sku": ["sku", "sifra", "sifra artikla", "sifra proizvoda", "artikal sifra", "item code", "code"],
    "ean": ["ean", "ean kod", "ean13", "ean 13", "gtin", "barkod", "barcode", "gtin upc ean or isbn"],
    "model": ["model", "mpn", "oznaka modela", "part number", "manufacturer part number"],
    "vpc": ["vpc", "redovna vpc", "vpc eur", "veleprodajna cijena", "nabavna cijena", "wholesale price"],
    "promo_vpc": ["akcija vpc", "akcijska vpc", "akc vpc", "vpc akcija", "promo vpc", "vpc promo", "akcijska veleprodajna cijena"],
    "brand": ["brand", "brands", "brend", "proizvodac", "marka", "manufacturer"],
    "category": ["kategorija", "kategorije", "category", "categories", "grupa"],
    "source_url": ["url", "source url", "izvor", "proizvodacki url", "product url", "link"],
    "images": ["slike", "slika", "images", "image url", "url slike"],
    "short": ["kratki opis", "short description"], "description": ["dugi opis", "opis", "description"],
}


def suggest_mapping(headers):
    result = {}
    for key, aliases in ALIASES.items():
        matches = [h for h in headers if norm(h) in aliases]
        if not matches and key == "promo_vpc":
            matches = [h for h in headers if ("vpc" in norm(h).split() or "veleprodaj" in norm(h)) and re.search(r"akc|promo", norm(h))]
        if not matches and key == "vpc":
            matches = [h for h in headers if "vpc" in norm(h).split() and not re.search(r"akc|promo", norm(h))]
        if len(matches) == 1:
            result[key] = matches[0]
    return result


def cell_text(value, number_format=""):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int) and re.fullmatch(r"0{2,}", number_format):
        return str(value).zfill(len(number_format))
    return str(value).strip()


def sheet_names(path):
    suffix = Path(path).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            return book.sheetnames
        finally:
            book.close()
    if suffix == ".xls":
        import xlrd
        with xlrd.open_workbook(path, on_demand=True) as book:
            return book.sheet_names()
    if suffix in (".csv", ".tsv"):
        return ["CSV"]
    raise ValidationError("Podržane datoteke: .xlsx, .xlsm, .xls, .csv i .tsv.")


def read_table(path, sheet=None, header_row=1):
    if not 1 <= header_row <= 1000:
        raise ValidationError("Redak zaglavlja mora biti između 1 i 1000.")
    suffix = Path(path).suffix.lower()
    rows = []
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        formulas = load_workbook(path, read_only=True, data_only=False)
        try:
            ws = book[sheet or book.sheetnames[0]]
            fw = formulas[ws.title]
            for cells, raw_cells in zip(ws.iter_rows(min_row=header_row), fw.iter_rows(min_row=header_row)):
                values = ["#FORMULA_BEZ_REZULTATA" if raw.data_type == "f" and c.value is None else cell_text(c.value, c.number_format) for c, raw in zip(cells, raw_cells)]
                rows.append(values)
                if len(rows) > 50001:
                    raise ValidationError("Najviše 50.000 redaka po listu. Podijelite tablicu.")
        finally:
            book.close()
            formulas.close()
    elif suffix == ".xls":
        import xlrd
        with xlrd.open_workbook(path, on_demand=True, formatting_info=True) as book:
            ws = book.sheet_by_name(sheet) if sheet else book.sheet_by_index(0)
            for row in range(header_row - 1, ws.nrows):
                values = []
                for col in range(ws.ncols):
                    cell = ws.cell(row, col)
                    fmt = book.format_map[book.xf_list[ws.cell_xf_index(row, col)].format_key].format_str
                    values.append("#EXCEL_ERROR" if cell.ctype == xlrd.XL_CELL_ERROR else cell_text(cell.value, fmt))
                rows.append(values)
    else:
        text = Path(path).read_text(encoding="utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:10000], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))[header_row - 1:]
    if not rows:
        raise ValidationError("Odabrani list nema podataka.")
    last = max((i for i, v in enumerate(rows[0]) if v), default=-1) + 1
    if not last:
        raise ValidationError("Odabrani redak zaglavlja je prazan.")
    headers = [v or f"Stupac {i+1}" for i, v in enumerate(rows[0][:last])]
    if len(set(headers)) != len(headers):
        raise ValidationError("Zaglavlje sadrži jednake nazive stupaca. Preimenujte ih radi nedvosmislenog mapiranja.")
    data = [(header_row + i, dict(zip(headers, row + [""] * max(0, len(headers) - len(row))))) for i, row in enumerate(rows[1:], 1) if any(row)]
    if len(data) > 50000:
        raise ValidationError("Najviše 50.000 redaka po listu.")
    return headers, data


def money(value):
    s = str("" if value is None else value).strip().replace("\u00a0", "").replace(" ", "")
    s = re.sub(r"(?i)EUR|€", "", s)
    if not s:
        return None
    if not re.fullmatch(r"\+?\d+(?:[.,]\d+)*", s):
        raise ValidationError(f"Neispravna cijena: {str(value)[:60]}")
    if "," in s and "." in s:
        dec, group = (",", ".") if s.rfind(",") > s.rfind(".") else (".", ",")
        whole, fraction = s.rsplit(dec, 1)
        if not re.fullmatch(r"\+?\d{1,3}(?:" + re.escape(group) + r"\d{3})+", whole) or len(fraction) > 2:
            raise ValidationError("Nejasni decimalni/tisućni razdjelnici cijene.")
        s = whole.replace(group, "") + "." + fraction
    elif "," in s or "." in s:
        sep = "," if "," in s else "."
        if s.count(sep) != 1 or len(s.split(sep)[1]) > 2:
            raise ValidationError("Nejasna cijena (npr. 1.250): koristite 1250 ili 1250,00; najviše 2 decimale.")
        s = s.replace(",", ".")
    try:
        result = Decimal(s)
    except InvalidOperation:
        raise ValidationError("Neispravna cijena.")
    if not result.is_finite() or not Decimal("0") < result <= Decimal("99999999"):
        raise ValidationError("Cijena mora biti pozitivna i manja od 100 milijuna EUR.")
    return result


@dataclass
class Settings:
    vat: str = "25"
    input_gross: bool = False
    price_mode: str = "preferred"
    model: str = "gpt-6-astra"
    domains: str = ""
    categories: str = ""
    limit: int = 20
    max_output_tokens: int = 6000
    batch_size: int = 100
    delimiter: str = ","
    export_mode: str = "all"
    default_category: str = ""

    def validate(self):
        try:
            vat = Decimal(self.vat.replace(",", "."))
        except InvalidOperation:
            raise ValidationError("PDV mora biti broj.")
        if not vat.is_finite() or not 0 <= vat <= 100:
            raise ValidationError("PDV mora biti između 0 i 100%.")
        if self.price_mode not in ("preferred", "sale"):
            raise ValidationError("Nepoznat način cijene.")
        if self.export_mode not in ("all", "new", "update", "split"):
            raise ValidationError("Nepoznat način izvoza.")
        if not 1 <= self.limit <= 5000 or not 1 <= self.batch_size <= 1000:
            raise ValidationError("Limit istraživanja: 1–5000. Veličina CSV-a: 1–1000.")
        if self.delimiter not in (",", ";"):
            raise ValidationError("CSV razdjelnik mora biti zarez ili točka-zarez.")
        return vat


def calculate_prices(vpc, promo, settings):
    vat = settings.validate()
    # A nonempty invalid promo must never silently fall back to another price.
    promo_amount = money(promo)
    base = money(vpc) if (settings.price_mode == "sale" or promo_amount is None) else None
    factor = Decimal(1) if settings.input_gross else Decimal(1) + vat / Decimal(100)
    gross = lambda n: str((n * factor).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
    if promo_amount is None and base is None:
        raise ValidationError("Nedostaje Akcija VPC i VPC.")
    if settings.price_mode == "sale" and promo_amount is not None and base is not None:
        regular, sale = gross(base), gross(promo_amount)
        if Decimal(sale) >= Decimal(regular):
            raise ValidationError("Akcijska cijena mora biti manja od redovne i nakon zaokruživanja.")
        return regular, sale, "VPC i Akcija VPC"
    return gross(promo_amount if promo_amount is not None else base), "", "Akcija VPC" if promo_amount is not None else "VPC"


def valid_ean(value):
    if not value:
        return True
    if not re.fullmatch(r"\d{8}|\d{12}|\d{13}|\d{14}", value):
        return False
    digits = list(map(int, value))
    return (sum(n * (3 if i % 2 == 0 else 1) for i, n in enumerate(reversed(digits[:-1]))) + digits[-1]) % 10 == 0


def url_shape(url):
    try:
        p = urlsplit(url)
        return p.scheme in ("https", "http") and bool(p.hostname) and not p.username and not p.password and not any(c.isspace() for c in url)
    except ValueError:
        return False


def parse_images(text):
    return list(dict.fromkeys(s.strip() for s in re.split(r"[\n|;]|,\s*(?=https?://)", text or "") if s.strip()))


@dataclass
class Product:
    uid: str = field(default_factory=lambda: uuid4().hex)
    origin: str = ""
    row: int = 0
    name: str = ""
    sku: str = ""
    ean: str = ""
    model: str = ""
    vpc: str = ""
    promo_vpc: str = ""
    regular: str = ""
    sale: str = ""
    price_source: str = ""
    brand: str = ""
    category: str = ""
    source_url: str = ""
    images: list = field(default_factory=list)
    short: str = ""
    description: str = ""
    specs: list = field(default_factory=list)
    evidence: str = ""
    researched_at: str = ""
    checked_images: list = field(default_factory=list)
    image_checked_at: str = ""
    status: str = "Čeka obradu"
    error: str = ""
    approval: str = ""
    notes: list = field(default_factory=list)
    usage: dict = field(default_factory=dict)

    def fingerprint(self):
        keys = ["name", "sku", "ean", "model", "regular", "sale", "brand", "category", "source_url", "images", "short", "description", "specs"]
        return hashlib.sha256(json.dumps({k: getattr(self, k) for k in keys}, ensure_ascii=False, sort_keys=True).encode()).hexdigest()

    def approve(self):
        problems = validate_product(self)
        if problems:
            raise ValidationError("\n".join(problems))
        self.approval = self.fingerprint()
        self.status, self.error = "Potvrđeno", ""

    @property
    def approved(self):
        return bool(self.approval) and self.approval == self.fingerprint()


def make_products(path, sheet, data, mapping, settings):
    if not mapping.get("name"):
        raise ValidationError("Odaberite stupac naziva proizvoda.")
    if not mapping.get("vpc") and not mapping.get("promo_vpc"):
        raise ValidationError("Odaberite barem jedan VPC stupac.")
    chosen = [v for v in mapping.values() if v]
    if len(chosen) != len(set(chosen)):
        raise ValidationError("Jedan stupac ne može biti mapiran na više polja.")
    products = []
    for row_no, values in data:
        fields = {key: values.get(column, "") for key, column in mapping.items() if key in FIELDS and column}
        fields["images"] = parse_images(fields.get("images", ""))
        p = Product(origin=f"{Path(path).name} / {sheet}", row=row_no, **fields)
        p.category = p.category or settings.default_category
        try:
            if any(str(v).startswith(("#FORMULA_BEZ_REZULTATA", "#EXCEL_ERROR", "#VALUE!", "#REF!", "#DIV/0!", "#N/A")) for k, v in fields.items() if k != "images"):
                raise ValidationError("Mapirani Excel podatak sadrži grešku ili formulu bez spremljenog izračuna.")
            if not p.name or p.name.startswith("#"):
                raise ValidationError("Nedostaje ispravan naziv proizvoda.")
            if not valid_ean(p.ean):
                raise ValidationError("Neispravan EAN/GTIN ili kontrolna znamenka.")
            p.regular, p.sale, p.price_source = calculate_prices(p.vpc, p.promo_vpc, settings)
            if not p.sku:
                p.sku = f"EAN-{p.ean}" if p.ean else "EP-" + hashlib.sha256(norm(p.name).encode()).hexdigest()[:14].upper()
                p.notes.append("SKU je automatski generiran; za postojeći proizvod provjerite originalni SKU/ID.")
        except ValidationError as exc:
            p.status, p.error = "Preskočeno", str(exc)
        products.append(p)
    mark_duplicates(products)
    return products


def mark_duplicates(products):
    counts = {key: Counter(getattr(p, key).casefold() for p in products if getattr(p, key)) for key in ("sku", "ean")}
    for p in products:
        for key, counter in counts.items():
            if getattr(p, key) and counter[getattr(p, key).casefold()] > 1:
                p.status = "Preskočeno"
                p.error = f"Duplikat {key.upper()} u učitanim redcima; nijedan duplikat nije odabran automatski."
                p.approval = ""


def validate_product(p, require_approval=False):
    errors = []
    for key, label in [("name", "naziv"), ("sku", "SKU"), ("brand", "brend"), ("category", "kategorija"), ("short", "kratki opis"), ("description", "dugi opis")]:
        v = getattr(p, key)
        if not v.strip():
            errors.append(f"Nedostaje {label}.")
        if v.lstrip().startswith(("=", "+", "-", "@")) or any(ord(c) < 32 and c not in "\n\r\t" for c in v):
            errors.append(f"Nedopušten početak ili kontrolni znak u polju {label}.")
    if not valid_ean(p.ean):
        errors.append("Neispravan EAN/GTIN.")
    try:
        regular = money(p.regular)
        sale = money(p.sale)
        if regular is None or (sale is not None and sale >= regular):
            errors.append("Cijena nedostaje ili akcijska nije manja od redovne.")
    except ValidationError as exc:
        errors.append(str(exc))
    if p.source_url and not url_shape(p.source_url):
        errors.append("Neispravan URL izvora.")
    if not p.images or any(not url_shape(x) or "," in x for x in p.images):
        errors.append("Nedostaju ispravni izravni URL-ovi slika (zarez kodirajte kao %2C).")
    if set(p.images) != set(p.checked_images) or not p.image_checked_at:
        errors.append("URL-ovi slika još nisu provjereni.")
    elif p.image_checked_at:
        try:
            if (datetime.now(timezone.utc) - datetime.fromisoformat(p.image_checked_at)).total_seconds() > 86400:
                errors.append("Provjera slika je starija od 24 sata; ponovno provjerite slike.")
        except (ValueError, TypeError):
            errors.append("Neispravno vrijeme provjere slika.")
    if require_approval and not p.approved:
        errors.append("Proizvod nije potvrđen nakon zadnje izmjene.")
    return errors


def html_description(p):
    paragraphs = "".join(f"<p>{html.escape(part).replace(chr(10), '<br>')}</p>" for part in p.description.split("\n\n") if part.strip())
    if p.specs:
        rows = "".join(f"<tr><th>{html.escape(str(s['name']))}</th><td>{html.escape(str(s['value']))}</td></tr>" for s in p.specs)
        paragraphs += f"<h2>Specifikacije</h2><table><tbody>{rows}</tbody></table>"
    return paragraphs


@contextmanager
def database(path, **kwargs):
    connection = sqlite3.connect(path, timeout=15, **kwargs)
    try:
        with connection:
            yield connection
    finally:
        connection.close()


class Store:
    def __init__(self, path):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with self.connect() as db:
            db.execute("CREATE TABLE IF NOT EXISTS products (uid TEXT PRIMARY KEY, body TEXT NOT NULL)")
            db.execute("CREATE TABLE IF NOT EXISTS settings (name TEXT PRIMARY KEY, body TEXT NOT NULL)")

    def connect(self):
        return database(self.path)

    def put(self, p):
        with self.connect() as db:
            db.execute("INSERT OR REPLACE INTO products VALUES (?,?)", (p.uid, json.dumps(asdict(p), ensure_ascii=False)))

    def load(self):
        with self.connect() as db:
            return [Product(**json.loads(row[0])) for row in db.execute("SELECT body FROM products ORDER BY rowid")]

    def replace(self, products):
        with self.connect() as db:
            db.execute("DELETE FROM products")
            db.executemany("INSERT INTO products VALUES (?,?)", [(p.uid, json.dumps(asdict(p), ensure_ascii=False)) for p in products])

    def get_setting(self, name, default=None):
        with self.connect() as db:
            row = db.execute("SELECT body FROM settings WHERE name=?", (name,)).fetchone()
        return json.loads(row[0]) if row else default

    def set_setting(self, name, value):
        with self.connect() as db:
            db.execute("INSERT OR REPLACE INTO settings VALUES (?,?)", (name, json.dumps(value, ensure_ascii=False)))

    def backup(self, target):
        with self.connect() as source, database(target) as destination:
            source.backup(destination)


class Catalog:
    def __init__(self, records):
        self.index = {"SKU": defaultdict(list), "GTIN, UPC, EAN, or ISBN": defaultdict(list)}
        for record in records:
            for field, index in self.index.items():
                key = record.get(field, "").strip().casefold()
                if key:
                    index[key].append(record)

    @classmethod
    def read(cls, path):
        headers, rows = read_table(path)
        if "SKU" not in headers and "GTIN, UPC, EAN, or ISBN" not in headers:
            raise ValidationError("Katalog mora biti WooCommerce CSV izvoz sa SKU ili GTIN stupcem.")
        return cls([r for _, r in rows])

    def match(self, p):
        matches = []
        for key, value in [("SKU", p.sku), ("GTIN, UPC, EAN, or ISBN", p.ean)]:
            found = self.index[key].get(value.casefold(), []) if value else []
            if len(found) > 1:
                raise ValidationError(f"Postojeći katalog ima više proizvoda za {key}.")
            matches.extend(found)
        distinct = {json.dumps(m, sort_keys=True): m for m in matches}
        if len(distinct) > 1:
            raise ValidationError("SKU i EAN pokazuju na različite postojeće proizvode.")
        if not distinct:
            return None
        match = next(iter(distinct.values()))
        existing_ean = match.get("GTIN, UPC, EAN, or ISBN", "").strip()
        if p.ean and existing_ean and p.ean != existing_ean:
            raise ValidationError("EAN proizvoda proturječi EAN-u postojećeg SKU-a.")
        if match.get("Type", "simple") != "simple":
            raise ValidationError("Postojeći proizvod nije simple; varijacije nisu podržane u ovoj verziji.")
        return match


def export_products(products, target, settings, catalog=None):
    settings.validate()
    if settings.export_mode != "all" and catalog is None:
        raise ValidationError("Za ovaj način učitajte postojeći WooCommerce CSV katalog radi provjere SKU/ID-a.")
    # New immutable run directory: a failed export never replaces an earlier delivery.
    folder = Path(target) / (datetime.now().strftime("izvoz-%Y%m%d-%H%M%S-") + uuid4().hex[:6])
    folder.mkdir(parents=True, exist_ok=False)
    groups = defaultdict(list)
    skipped, provenance = [], []
    duplicate_sku = Counter(p.sku.casefold() for p in products if p.sku)
    duplicate_ean = Counter(p.ean for p in products if p.ean)
    targets = set()
    for p in products:
        try:
            problems = validate_product(p, require_approval=True)
            if p.error:
                problems.append(p.error)
            if duplicate_sku[p.sku.casefold()] > 1 or (p.ean and duplicate_ean[p.ean] > 1):
                problems.append("Duplikat SKU/EAN u radnom katalogu.")
            if problems:
                raise ValidationError(" ".join(problems))
            existing = catalog.match(p) if catalog else None
            if settings.export_mode == "new" and existing:
                raise ValidationError("Postojeći proizvod preskočen u načinu Samo novi.")
            if settings.export_mode == "update" and not existing:
                raise ValidationError("Proizvod ne postoji u učitanom WooCommerce katalogu.")
            update = bool(existing and settings.export_mode in ("update", "split"))
            if existing and not update:
                raise ValidationError("Postojeći proizvod preskočen pri pripremi novih proizvoda.")
            record = {"SKU": p.sku, "Name": p.name, "Short description": "<p>" + html.escape(p.short).replace("\n", "<br>") + "</p>", "Description": html_description(p), "Regular price": p.regular,
                      "Categories": p.category, "Brands": p.brand.replace(",", "\\,"), "Images": ", ".join(p.images)}
            if update:
                record["SKU"] = existing.get("SKU", "")
                if existing.get("ID", ""):
                    if not str(existing["ID"]).isdigit():
                        raise ValidationError("Neispravan ID u postojećem katalogu.")
                    record["ID"] = existing["ID"]
                if not record["SKU"] and not record.get("ID"):
                    raise ValidationError("Nema postojećeg SKU/ID za ažuriranje.")
                identity = record.get("ID") or record["SKU"].casefold()
                if identity in targets:
                    raise ValidationError("Više ulaznih redaka pokušava ažurirati isti proizvod.")
                targets.add(identity)
            else:
                record.update({"Type": "simple", "Published": "-1"})
                if p.ean:
                    record["GTIN, UPC, EAN, or ISBN"] = p.ean
            if p.sale:
                record["Sale price"] = p.sale
            # Missing values are omitted from update columns, preserving existing data.
            group = ("azuriranje" if update else "novi", tuple(sorted(record)))
            groups[group].append(record)
            provenance.append({"SKU": p.sku, "Naziv": p.name, "Izvor": p.source_url, "Provjereno": p.researched_at, "Ulaz": p.origin, "Redak": p.row, "Osnova cijene": p.price_source, "Cijena s PDV": p.regular, "Akcijska s PDV": p.sale})
        except ValidationError as exc:
            skipped.append(f"{p.origin} | redak {p.row} | {p.name} | {exc}")
    files = []
    for (kind, fields), records in groups.items():
        for start in range(0, len(records), settings.batch_size):
            path = folder / f"{kind}-{len(files)+1:03d}.csv"
            with path.open("w", newline="", encoding="utf-8-sig") as stream:
                writer = csv.DictWriter(stream, fieldnames=fields, delimiter=settings.delimiter, lineterminator="\r\n")
                writer.writeheader()
                writer.writerows(records[start:start + settings.batch_size])
            with path.open(newline="", encoding="utf-8-sig") as stream:
                reread = list(csv.DictReader(stream, delimiter=settings.delimiter))
            if reread != records[start:start + settings.batch_size]:
                raise ValidationError("Provjera spremljenog CSV-a nije prošla.")
            files.append(path.name)
    if provenance:
        with (folder / "izvori.csv").open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=list(provenance[0]), delimiter=";")
            writer.writeheader()
            for record in provenance:
                writer.writerow({k: ("'" + str(v) if str(v).lstrip().startswith(("=", "+", "-", "@")) else v) for k, v in record.items()})
    summary = {"ulaz": len(products), "izvezeno": len(provenance), "preskoceno": len(skipped), "csv": files, "vrijeme": now(), "postavke": asdict(settings)}
    (folder / "izvjestaj.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    (folder / "log.txt").write_text(f"{now()} | Ulaz: {len(products)} | Izvezeno: {len(provenance)} | Preskočeno: {len(skipped)}\n" + "\n".join(skipped), encoding="utf-8-sig")
    (folder / "UPUTE.txt").write_text("Uvozite samo novi-*.csv i azuriranje-*.csv. izvori.csv je izvještaj.\n"
        "WooCommerce > Proizvodi > Uvoz. Za azuriranje-*.csv uključite Ažuriraj postojeće proizvode; za novi-*.csv isključite.\n"
        "Novi proizvodi imaju status Nacrt. Ažuriranje zadržava postojeći status i zalihe.\n"
        "Sve cijene su EUR s PDV-om. WooCommerce postavka unosa cijena mora biti S uključenim porezom.\n"
        "Prazna akcijska cijena se ne izvozi: postojeća akcija se time NE uklanja.\n"
        "Provjerite mapiranje Brands i GTIN prema verziji trgovine. Razdjelnik CSV-a: " + settings.delimiter + "\n"
        "Krenite s malim probnim uvozom. Ovaj program ne pristupa WordPressu.\n", encoding="utf-8-sig")
    return folder, summary

